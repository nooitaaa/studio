import argparse
import re
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image


ROOT = Path(__file__).resolve().parent
HTML_PATH = ROOT / "index.html"
WORKBOOK_PATH = ROOT / "didascalie.xlsx"
ENTRY_PATTERN = re.compile(
    r"\{\s*title:\s*'(?P<title>(?:\\.|[^'\\])*)',\s*"
    r"caption:\s*'(?P<caption>(?:\\.|[^'\\])*)',\s*"
    r"(?:project:\s*'(?P<project>(?:\\.|[^'\\])*)',\s*)?"
    r"files:\s*\[(?P<files>[^\]]*)\]\s*\}"
)


def decode_js_string(value):
    return value.replace("\\'", "'").replace("\\\\", "\\")


def encode_js_string(value):
    escaped = value.replace("\\", "\\\\").replace("'", "\\'").replace("\r", "\\r").replace("\n", "\\n")
    return f"'{escaped}'"


def read_projects():
    source = HTML_PATH.read_text(encoding="utf-8")
    projects = []
    for match in ENTRY_PATTERN.finditer(source):
        files = [decode_js_string(name) for name in re.findall(r"'((?:\\.|[^'\\])*)'", match.group("files"))]
        projects.append({
            "title": decode_js_string(match.group("title")),
            "caption": decode_js_string(match.group("caption")),
            "project": decode_js_string(match.group("project") or match.group("title")),
            "files": files,
        })
    if not projects:
        raise ValueError("Non trovo i progetti in index.html.")
    return source, projects


def add_preview(sheet, cell, image_path, temp_dir, cache):
    if image_path not in cache:
        thumbnail_path = Path(temp_dir) / f"{len(cache)}.jpg"
        with Image.open(image_path) as image:
            image.thumbnail((150, 108))
            if image.mode != "RGB":
                image = image.convert("RGB")
            image.save(thumbnail_path, "JPEG", quality=78, optimize=True)
        cache[image_path] = thumbnail_path

    preview = ExcelImage(str(cache[image_path]))
    sheet.add_image(preview, cell)


def create_workbook(projects):
    workbook = Workbook()
    captions_sheet = workbook.active
    captions_sheet.title = "Didascalie"
    captions_sheet.sheet_view.showGridLines = False
    captions_sheet.merge_cells("A1:D1")
    captions_sheet["A1"] = "DIDASCALIE DEL PORTFOLIO"
    captions_sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    captions_sheet["A1"].fill = PatternFill("solid", fgColor="183B3B")
    captions_sheet["A1"].alignment = Alignment(vertical="center")
    captions_sheet.row_dimensions[1].height = 34
    captions_sheet.merge_cells("A2:D2")
    captions_sheet["A2"] = (
        "Modifica la colonna D. Per applicare le modifiche al sito, salva il file e avvia "
        "python3 aggiorna_didascalie.py dal Terminale nella cartella del sito."
    )
    captions_sheet["A2"].font = Font(name="Aptos", size=10, color="334155")
    captions_sheet["A2"].fill = PatternFill("solid", fgColor="E8F0ED")
    captions_sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    captions_sheet.row_dimensions[2].height = 34
    headers = ["Anteprima", "Progetto", "Immagini associate", "Didascalia modificabile"]
    for column, header in enumerate(headers, start=1):
        cell = captions_sheet.cell(row=3, column=column, value=header)
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="356A61")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    captions_sheet.column_dimensions["A"].width = 23
    captions_sheet.column_dimensions["B"].width = 29
    captions_sheet.column_dimensions["C"].width = 66
    captions_sheet.column_dimensions["D"].width = 82
    captions_sheet.row_dimensions[3].height = 30
    captions_sheet.freeze_panes = "B4"

    images_sheet = workbook.create_sheet("Immagini")
    images_sheet.sheet_view.showGridLines = False
    image_headers = ["Anteprima", "Progetto", "File immagine", "Didascalia progetto"]
    for column, header in enumerate(image_headers, start=1):
        cell = images_sheet.cell(row=1, column=column, value=header)
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="356A61")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    images_sheet.column_dimensions["A"].width = 23
    images_sheet.column_dimensions["B"].width = 29
    images_sheet.column_dimensions["C"].width = 40
    images_sheet.column_dimensions["D"].width = 82
    images_sheet.row_dimensions[1].height = 30
    images_sheet.freeze_panes = "B2"

    preview_cache = {}
    with tempfile.TemporaryDirectory(prefix="didascalie-") as temp_dir:
        for project_index, project in enumerate(projects, start=4):
            first_image = ROOT / "images" / project["files"][0]
            if not first_image.is_file():
                raise FileNotFoundError(f"Immagine non trovata: {first_image}")
            captions_sheet.cell(project_index, 2, project["title"])
            captions_sheet.cell(project_index, 3, ", ".join(project["files"]))
            caption_cell = captions_sheet.cell(project_index, 4, project["caption"])
            caption_cell.fill = PatternFill("solid", fgColor="FFF4D6")
            caption_cell.comment = None
            for cell in captions_sheet[project_index][1:]:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.font = Font(name="Aptos", size=10, color="183B3B")
            captions_sheet.row_dimensions[project_index].height = 88
            add_preview(captions_sheet, f"A{project_index}", first_image, temp_dir, preview_cache)

            for filename in project["files"]:
                image_path = ROOT / "images" / filename
                if not image_path.is_file():
                    raise FileNotFoundError(f"Immagine non trovata: {image_path}")
                row = images_sheet.max_row + 1
                images_sheet.cell(row, 2, project["title"])
                images_sheet.cell(row, 3, filename)
                images_sheet.cell(row, 4, f"='Didascalie'!D{project_index}")
                for cell in images_sheet[row][1:]:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    cell.font = Font(name="Aptos", size=10, color="183B3B")
                images_sheet.cell(row, 4).fill = PatternFill("solid", fgColor="E8F0ED")
                images_sheet.row_dimensions[row].height = 88
                add_preview(images_sheet, f"A{row}", image_path, temp_dir, preview_cache)

        captions_sheet.auto_filter.ref = f"A3:D{len(projects) + 3}"
        images_sheet.auto_filter.ref = f"A1:D{images_sheet.max_row}"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(WORKBOOK_PATH)


def apply_workbook():
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError("Non trovo didascalie.xlsx. Generala prima con --create.")

    source, projects = read_projects()
    workbook = load_workbook(WORKBOOK_PATH, data_only=False)
    sheet = workbook["Didascalie"]
    captions = {}
    for row in range(4, sheet.max_row + 1):
        title = sheet.cell(row, 2).value
        caption = sheet.cell(row, 4).value
        if title:
            if caption is None or not str(caption).strip():
                raise ValueError(f"La didascalia del progetto '{title}' è vuota.")
            captions[str(title)] = str(caption)

    matches = {project["title"]: project for project in projects}
    missing = sorted(set(matches) - set(captions))
    unknown = sorted(set(captions) - set(matches))
    if missing or unknown:
        raise ValueError(f"Progetti mancanti: {missing}; progetti non riconosciuti: {unknown}.")

    replacements = []
    for match in ENTRY_PATTERN.finditer(source):
        title = decode_js_string(match.group("title"))
        current_caption = decode_js_string(match.group("caption"))
        if title in captions and captions[title] != current_caption:
            replacements.append((match.start("caption") - 1, match.end("caption") + 1, encode_js_string(captions[title])))
    for start, end, replacement in reversed(replacements):
        source = source[:start] + replacement + source[end:]

    HTML_PATH.write_text(source, encoding="utf-8")
    print(f"Aggiornate {len(replacements)} didascalie in index.html.")


def main():
    parser = argparse.ArgumentParser(description="Crea o applica le didascalie del portfolio.")
    parser.add_argument("--create", action="store_true", help="Genera didascalie.xlsx con le immagini del portfolio.")
    args = parser.parse_args()

    try:
        if args.create:
            _, projects = read_projects()
            create_workbook(projects)
            print(f"Creato {WORKBOOK_PATH.name}: {len(projects)} progetti, immagini incorporate.")
        else:
            apply_workbook()
    except Exception as error:
        print(f"Errore: {error}", file=sys.stderr)
        raise SystemExit(1)


if __name__ == "__main__":
    main()